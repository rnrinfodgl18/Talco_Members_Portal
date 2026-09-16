import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.importer.matcher import resolve
from app.importer.models import Tannery as ImportTannery
from app.importer.normalize import nkey, split_party
from app.models import AuditLog, Party, PartyAlias, Tannery, TanneryPartyLink
from app.services.alias_resolver import approved_aliases


_REPORT_DATE = re.compile(r"\bFor\s+(\d{1,2})-([A-Za-z]{3})-(\d{4})\b", re.IGNORECASE)
_MONTHS = {name: number for number, name in enumerate(
    ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1)}


def _amount(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid amount {value!r} in the opening balance file") from exc


def parse_opening_workbook(content: bytes) -> tuple[date, list[dict], Decimal, Decimal]:
    if not content:
        raise ValueError("Choose an Excel opening balance file")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("The opening balance file must be a readable .xlsx workbook") from exc
    try:
        sheet = workbook.active
        report_date = None
        header_row = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = list(row)
            for value in values:
                match = _REPORT_DATE.search(str(value or ""))
                if match:
                    report_date = date(int(match.group(3)), _MONTHS[match.group(2).upper()], int(match.group(1)))
            if str(values[0] or "").strip().lower() == "particulars":
                header_row = row_number
                break
        if report_date is None:
            raise ValueError("Could not find the Tally report date (for example: For 1-Apr-2026)")
        if header_row is None:
            raise ValueError("Could not find Particulars, Debit and Credit columns")

        entries: list[dict] = []
        grand_debit = grand_credit = None
        seen: set[str] = set()
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            name = str(row[0] or "").strip()
            if not name:
                continue
            debit = _amount(row[1] if len(row) > 1 else None)
            credit = _amount(row[2] if len(row) > 2 else None)
            if name.lower() == "grand total":
                grand_debit, grand_credit = debit, credit
                break
            if debit < 0 or credit < 0 or (debit and credit):
                raise ValueError(f"{name}: enter a positive amount in either Debit or Credit, not both")
            key = nkey(name)
            if not key or key in seen:
                raise ValueError(f"Duplicate or invalid ledger name: {name}")
            seen.add(key)
            entries.append({"name": name, "debit": debit, "credit": credit,
                            "amount": debit - credit, "side": "Dr" if debit >= credit else "Cr"})
        if not entries:
            raise ValueError("No ledger balances were found in the workbook")
        debit_total = sum((row["debit"] for row in entries), Decimal("0"))
        credit_total = sum((row["credit"] for row in entries), Decimal("0"))
        if grand_debit is not None and (debit_total != grand_debit or credit_total != grand_credit):
            raise ValueError("Ledger amounts do not reconcile with the Grand Total row")
        return report_date, entries, debit_total, credit_total
    finally:
        workbook.close()


def _context(session: Session):
    db_tanneries = session.scalars(select(Tannery).order_by(Tannery.sno)).all()
    tanneries = [ImportTannery(sno=row.sno, name=row.name, pump_house=row.pump_house) for row in db_tanneries]
    by_sno = {row.sno: row for row in db_tanneries}
    aliases = {row.normalized_key: row for row in session.scalars(select(PartyAlias)).all()
               if not row.excluded and row.revoked_at is None and row.party_id is not None}
    parties = {row.id: row for row in session.scalars(select(Party)).all()}
    parties_by_key = {row.normalized_key: row for row in parties.values()}
    return tanneries, by_sno, aliases, parties, parties_by_key


def preview_opening_balances(session: Session, filename: str, content: bytes) -> dict:
    report_date, entries, debit_total, credit_total = parse_opening_workbook(content)
    tanneries, by_sno, db_aliases, parties, parties_by_key = _context(session)
    accepted = approved_aliases(session)
    rows = []
    for index, entry in enumerate(entries, 1):
        parts = split_party(entry["name"])
        alias = db_aliases.get(nkey(parts.cleaned))
        party = parties.get(alias.party_id) if alias else parties_by_key.get(nkey(parts.left))
        match = resolve(entry["name"], tanneries, accepted)
        tannery = by_sno.get(match.tannery.sno) if match.tannery else None
        if party and tannery is None:
            link = session.scalar(select(TanneryPartyLink).where(
                TanneryPartyLink.party_id == party.id, TanneryPartyLink.valid_to.is_(None)))
            tannery = session.get(Tannery, link.tannery_id) if link else None
        current = party.opening_amount if party else Decimal("0")
        current_date = party.opening_date if party else None
        if tannery is None:
            action = "unmatched"
        elif party is None:
            action = "create"
        elif current == entry["amount"] and current_date == report_date:
            action = "unchanged"
        elif current_date is None and current == 0:
            action = "set"
        else:
            action = "overwrite"
        rows.append({"row": index, "ledger_name": entry["name"], "account_id": party.id if party else None,
                     "account_name": party.name if party else parts.left, "tannery_id": tannery.id if tannery else None,
                     "tannery_name": tannery.name if tannery else None, "match_method": match.method,
                     "current_amount": str(abs(current)), "current_side": "Cr" if current < 0 else "Dr",
                     "current_date": current_date.isoformat() if current_date else None,
                     "proposed_amount": str(abs(entry["amount"])), "proposed_side": entry["side"],
                     "action": action})
    return {"file_name": filename, "file_sha256": hashlib.sha256(content).hexdigest(),
            "opening_date": report_date.isoformat(), "row_count": len(rows),
            "debit_total": str(debit_total), "credit_total": str(credit_total),
            "create_count": sum(row["action"] == "create" for row in rows),
            "set_count": sum(row["action"] == "set" for row in rows),
            "overwrite_count": sum(row["action"] == "overwrite" for row in rows),
            "unchanged_count": sum(row["action"] == "unchanged" for row in rows),
            "unmatched_count": sum(row["action"] == "unmatched" for row in rows), "rows": rows}


def apply_opening_balances(session: Session, filename: str, content: bytes, approve_overwrite: bool,
                           reason: str, actor: str) -> dict:
    preview = preview_opening_balances(session, filename, content)
    if preview["unmatched_count"]:
        raise ValueError("Resolve every unmatched tannery before importing opening balances")
    if preview["overwrite_count"] and not approve_overwrite:
        raise ValueError("Review and approve the changed opening balances before overwrite")
    reason = reason.strip()
    if preview["overwrite_count"] and not reason:
        raise ValueError("Enter an approval reason for changed opening balances")

    report_date = date.fromisoformat(preview["opening_date"])
    _, entries, _, _ = parse_opening_workbook(content)
    entry_by_name = {row["name"]: row for row in entries}
    tanneries, by_sno, db_aliases, parties, parties_by_key = _context(session)
    accepted = approved_aliases(session)
    changed = 0
    for row in preview["rows"]:
        if row["action"] == "unchanged":
            continue
        entry = entry_by_name[row["ledger_name"]]
        parts = split_party(entry["name"])
        alias = db_aliases.get(nkey(parts.cleaned))
        party = parties.get(alias.party_id) if alias else parties_by_key.get(nkey(parts.left))
        match = resolve(entry["name"], tanneries, accepted)
        tannery = by_sno[match.tannery.sno]
        if party is None:
            party = Party(name=parts.left, normalized_key=nkey(parts.left))
            session.add(party)
            session.flush()
            parties[party.id] = party
            parties_by_key[party.normalized_key] = party
        alias_key = nkey(parts.cleaned)
        existing_alias = session.scalar(select(PartyAlias).where(PartyAlias.normalized_key == alias_key))
        if existing_alias is None:
            existing_alias = PartyAlias(party_id=party.id, tannery_id=tannery.id, normalized_key=alias_key,
                                        raw_name=parts.cleaned, excluded=False)
            session.add(existing_alias)
        elif existing_alias.party_id not in (None, party.id):
            raise ValueError(f"{entry['name']}: ledger alias is linked to another account")
        else:
            existing_alias.party_id, existing_alias.tannery_id, existing_alias.excluded, existing_alias.revoked_at = party.id, tannery.id, False, None
        active_links = session.scalars(select(TanneryPartyLink).where(
            TanneryPartyLink.party_id == party.id, TanneryPartyLink.valid_to.is_(None))).all()
        if not any(link.tannery_id == tannery.id for link in active_links):
            session.add(TanneryPartyLink(tannery_id=tannery.id, party_id=party.id,
                                         role="owner", valid_from=report_date))
        old = {"amount": str(party.opening_amount or 0),
               "date": party.opening_date.isoformat() if party.opening_date else None}
        party.opening_amount = entry["amount"]
        party.opening_date = report_date
        party.opening_note = f"Imported from {filename}; Tally balance as at {report_date.isoformat()}"
        session.flush()
        session.add(AuditLog(entity_type="ledger_account", entity_id=party.id,
                             action="opening_overwrite" if row["action"] == "overwrite" else "opening_import",
                             changed_by=actor, changes=json.dumps({"from": old,
                                 "to": {"amount": str(entry["amount"]), "date": report_date.isoformat()},
                                 "file": filename, "file_sha256": preview["file_sha256"],
                                 "reason": reason or "Initial opening balance import"})))
        changed += 1
    session.commit()
    return {**preview, "status": "applied", "updated_count": changed,
            "message": f"{changed} opening balance(s) applied; {preview['unchanged_count']} unchanged"}
