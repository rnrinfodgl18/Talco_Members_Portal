import re
from pathlib import Path

from openpyxl import load_workbook

from app.importer.models import Tannery


def _text(value: object) -> str | None:
    if value is None:
        return None
    return str(value).strip() or None


def _integer(value: object) -> int | None:
    if value is None or str(value).strip().upper() == "NIL":
        return None
    try:
        return int(float(str(value).replace(",", "")))
    except ValueError:
        return None


def _gstin(value: object) -> str | None:
    raw = _text(value)
    if not raw:
        return None
    match = re.search(r"[0-9A-Z\u0396]{15}", raw.upper())
    return match.group(0) if match else raw


def _sheet_number(sheet) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 4), values_only=True):
        for value in row:
            match = re.match(r"\s*(\d+)\s*[.]", str(value or ""))
            if match:
                return int(match.group(1))
    raise ValueError(f"Sheet {sheet.title!r}: missing tannery serial number")


def _fields(sheet) -> dict[str, object]:
    fields: dict[str, object] = {}
    for row in sheet.iter_rows(values_only=True):
        if len(row) >= 3 and isinstance(row[0], (int, float)) and row[1]:
            fields[str(row[1]).strip().upper()] = row[2]
    return fields


def _find(fields: dict[str, object], starts_with: str) -> object:
    return next((value for label, value in fields.items() if label.startswith(starts_with)), None)


def load_workbooks(paths: list[str | Path]) -> list[Tannery]:
    tanneries: list[Tannery] = []
    for path_value in paths:
        path = Path(path_value)
        pump_match = re.match(r"([A-E])\s+PUMP", path.name, re.IGNORECASE)
        if not pump_match:
            raise ValueError(f"Cannot determine pump house from {path.name}")
        pump = pump_match.group(1).upper()
        workbook = load_workbook(path, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            if sheet.max_row <= 1:
                continue
            fields = _fields(sheet)
            name = _text(_find(fields, "NAME OF THE TANNERY"))
            if not name:
                continue
            tanneries.append(Tannery(
                sno=_sheet_number(sheet), name=name, pump_house=pump,
                internal_id=_text(_find(fields, "TANNERY ID INTERNAL")),
                factory_id=_text(_find(fields, "TANNERY ID (")),
                tnpcb_user_id=_text(_find(fields, "TNPCB STATUS")),
                gps=_text(_find(fields, "TANNERY GPS")), gstin=_gstin(_find(fields, "GST NUMBER")),
                consent=_text(_find(fields, "TNPCB CONSENT DETAILS")),
                original_capacity=_text(_find(fields, "CAPACITY (ORIGINAL")),
                additional_capacity=_text(_find(fields, "CAPACITY ( ADDITIONAL")),
                original_shares=_integer(_find(fields, "SHARES HELD")),
                additional_shares=_integer(_find(fields, "SHARES ( ADDITIONAL")),
                attributes={label: _text(value) for label, value in fields.items()},
            ))
        workbook.close()
    serials = [item.sno for item in tanneries]
    if len(serials) != len(set(serials)):
        raise ValueError("Duplicate tannery serial number in master workbooks")
    return sorted(tanneries, key=lambda item: item.sno)


def load_master_directory(directory: str | Path) -> list[Tannery]:
    return load_workbooks(sorted(Path(directory).glob("[A-E] PUMP*.xlsx")))

