from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from app.importer.models import VoucherRow
from app.importer.parser import TallyParseError


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def _date(value, sheet: str, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise TallyParseError(f"{sheet} row {row}: Date is required or invalid")


def _money(value, sheet: str, row: int, label: str) -> Decimal:
    try:
        if value is None or _text(value) == "":
            raise InvalidOperation
        return abs(Decimal(str(value).replace(",", "").replace("₹", "").strip()))
    except InvalidOperation as exc:
        raise TallyParseError(f"{sheet} row {row}: {label} is required or invalid") from exc


def _key(value) -> str:
    return " ".join(_text(value).casefold().replace(".", "").split())


def _header_map(values) -> dict[str, int]:
    aliases = {
        "date": {"date"},
        "party": {"party", "party name", "ledger", "ledger name", "particulars"},
        "voucher_type": {"vch type", "voucher type", "type"},
        "voucher_no": {"vch no", "vch no.", "voucher no", "voucher number", "bill no", "bill number"},
        "base": {"taxable amount", "base amount", "amount", "credit"},
        "cgst": {"cgst", "cgst amount"},
        "sgst": {"sgst", "sgst amount"},
        "gross": {"invoice total", "gross amount", "total amount"},
        "bill_ref": {"against invoice", "against invoice number", "bill ref", "bill reference", "against ref"},
        "cancelled": {"cancelled", "is cancelled", "status"},
        "charge": {"charge head", "charge ledger", "service"},
    }
    found = {}
    for index, value in enumerate(values):
        key = _key(value)
        for name, choices in aliases.items():
            if key in choices and name not in found:
                found[name] = index
    return found


def _sheet_charge(ws, headers: dict[str, int]) -> tuple[str | None, Decimal | None]:
    title = " ".join(_text(ws.cell(row, 1).value) for row in range(1, min(ws.max_row, 12) + 1)).casefold()
    if "sludge" in title:
        return "Sludge disposal", Decimal("18")
    if "chrome" in title:
        return "Chrome water", Decimal("5")
    if "treatment" in title or "effluent" in title:
        return "Treatment charges", Decimal("5")
    return None, None


def _value(values, headers: dict[str, int], name: str):
    index = headers.get(name)
    return values[index] if index is not None and index < len(values) else None


def parse_excel(path: str | Path) -> list[VoucherRow]:
    source = Path(path)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise TallyParseError(f"Cannot read Excel workbook {source.name}: {exc}") from exc

    rows: list[VoucherRow] = []
    row_number = 0
    for ws in workbook.worksheets:
        header_row = None
        headers = {}
        for number, values in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True), 1):
            candidate = _header_map(values)
            if "date" in candidate and ("voucher_type" in candidate or "party" in candidate):
                header_row, headers = number, candidate
                break
        if header_row is None:
            continue
        sheet_charge, sheet_rate = _sheet_charge(ws, headers)
        for excel_row, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            voucher_type = _text(_value(values, headers, "voucher_type"))
            raw_date = _value(values, headers, "date")
            if not raw_date or not voucher_type or voucher_type.casefold() not in {"sales", "receipt", "bank receipt", "cash receipt"}:
                continue
            row_number += 1
            warnings = []
            party = _text(_value(values, headers, "party"))
            # Tally Sales ledger reports place the actual party in the column after the 'By' marker.
            party_index = headers.get("party")
            if voucher_type.casefold() == "sales" and party.casefold() in {"by", "to"} and party_index is not None:
                party = _text(values[party_index + 1] if party_index + 1 < len(values) else None)
            if not party:
                raise TallyParseError(f"{ws.title} row {excel_row}: Party / Ledger Name is required")
            voucher_no = _text(_value(values, headers, "voucher_no")) or None
            if not voucher_no:
                warnings.append("Voucher / bill number missing; a stable fallback reference will be generated")
            is_receipt = "receipt" in voucher_type.casefold()
            bill_ref = _text(_value(values, headers, "bill_ref")) or None
            allocations = ()
            if is_receipt:
                gross = _money(_value(values, headers, "gross") or _value(values, headers, "base"),
                               ws.title, excel_row, "Receipt amount")
                if bill_ref:
                    allocations = ((bill_ref, "Agst Ref", gross),)
                else:
                    warnings.append("Against invoice / bill reference missing; receipt will remain unallocated")
                base = cgst = sgst = rate = None
                charge = None
                voucher_type = "Receipt"
            else:
                charge = _text(_value(values, headers, "charge")) or sheet_charge
                rate = sheet_rate
                if not charge or rate is None:
                    raise TallyParseError(f"{ws.title} row {excel_row}: Charge Head is required")
                base = _money(_value(values, headers, "base"), ws.title, excel_row, "Taxable amount")
                raw_cgst, raw_sgst = _value(values, headers, "cgst"), _value(values, headers, "sgst")
                if raw_cgst not in (None, "") and raw_sgst not in (None, ""):
                    cgst = _money(raw_cgst, ws.title, excel_row, "CGST")
                    sgst = _money(raw_sgst, ws.title, excel_row, "SGST")
                else:
                    cgst = (base * rate / Decimal("200")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
                    sgst = cgst
                    warnings.append(f"CGST and SGST calculated from {rate}% charge-head rate")
                raw_gross = _value(values, headers, "gross")
                gross = _money(raw_gross, ws.title, excel_row, "Invoice total") if raw_gross not in (None, "") else base + cgst + sgst
                if raw_gross in (None, ""):
                    warnings.append("Invoice total calculated from taxable amount plus CGST and SGST")
                allocations = ((voucher_no, "New Ref", gross),) if voucher_no else ()
                voucher_type = "Sales"
            status = _text(_value(values, headers, "cancelled")).casefold()
            rows.append(VoucherRow(
                row_number=row_number, voucher_date=_date(raw_date, ws.title, excel_row),
                party=party, voucher_type=voucher_type, gstin=None, gross=gross,
                ledger_amounts=tuple(x for x in (base, cgst, sgst) if x is not None),
                base=base, cgst=cgst, sgst=sgst, gst_rate=rate, voucher_no=voucher_no,
                charge_ledger=charge, bill_allocations=allocations,
                is_cancelled=status in {"yes", "true", "cancelled", "canceled"},
                warnings=tuple(warnings), source_format="excel",
            ))
    if not rows:
        raise TallyParseError(f"No supported Sales or Receipt rows found in {source.name}")
    return rows
